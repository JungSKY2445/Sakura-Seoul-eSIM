"""
LINE eSIM Bot - メインアプリケーション
Amazon eSIM自動発行システム
"""
import io
import re
import hashlib
import secrets
import functools
import logging
import qrcode
import requests as http_requests
from flask import Flask, request, abort, jsonify, send_file, Response
from PIL import Image, ImageDraw, ImageFont

from linebot.v3 import WebhookHandler
from linebot.v3.exceptions import InvalidSignatureError
from linebot.v3.messaging import (
    Configuration,
    ApiClient,
    MessagingApi,
    MessagingApiBlob,
    ReplyMessageRequest,
    PushMessageRequest,
    TextMessage,
    ImageMessage,
)
from linebot.v3.webhooks import (
    MessageEvent,
    TextMessageContent,
    FollowEvent,
)

import config
import database as db
import messages as msg
import amazon_api

# Flask app
app = Flask(__name__)

# DB 자동 초기화 (서버 시작 시)
db.init_db()

# LINE SDK v3 설정
line_config = Configuration(access_token=config.LINE_CHANNEL_ACCESS_TOKEN)
handler = WebhookHandler(config.LINE_CHANNEL_SECRET)

# Amazon 주문번호 정규식 (예: 250-1234567-1234567)
AMAZON_ORDER_PATTERN = re.compile(r'^\d{3}-\d{7}-\d{7}$')

logger = logging.getLogger(__name__)


# ========== Admin 인증 ==========

def require_admin(f):
    """Admin API 키 인증 데코레이터"""
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if not config.ADMIN_API_KEY:
            logger.warning("ADMIN_API_KEY not set — admin access blocked")
            return jsonify({"error": "Admin access not configured"}), 403

        # Header: Authorization: Bearer <key>  또는  ?key=<key>
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            provided_key = auth[7:]
        else:
            provided_key = request.args.get("key", "")

        if not provided_key or not secrets.compare_digest(provided_key, config.ADMIN_API_KEY):
            logger.warning(f"Admin auth failed from {request.remote_addr}")
            return jsonify({"error": "Unauthorized"}), 401

        return f(*args, **kwargs)
    return decorated


# ========== Webhook Endpoint ==========

@app.route("/callback", methods=["POST"])
def callback():
    signature = request.headers.get("X-Line-Signature", "")
    body = request.get_data(as_text=True)
    app.logger.info(f"Request body: {body}")

    # LINE Verify 요청 (빈 events) 대응
    if not body or body == "{}":
        return "OK"

    try:
        handler.handle(body, signature)
    except InvalidSignatureError:
        app.logger.error("Invalid signature")
        abort(400)
    except Exception as e:
        app.logger.error(f"Error: {e}")
        return "OK"

    return "OK"


@app.route("/health", methods=["GET"])
def health():
    """기본 헬스체크 (인증 불필요)"""
    return {"status": "ok"}


@app.route("/qr/<token>", methods=["GET"])
def serve_qr_image(token):
    """QR코드 이미지 제공 (토큰 인증으로 외부 노출 방지)"""
    # 토큰으로 주문 조회
    order = db.get_order_by_qr_token(token)
    if not order or not order.get("qr_code_data"):
        abort(404)

    # QR코드 이미지 로컬 생성
    qr = qrcode.QRCode(version=None, box_size=10, border=4,
                        error_correction=qrcode.constants.ERROR_CORRECT_L)
    qr.add_data(order["qr_code_data"])
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return Response(buf.getvalue(), mimetype="image/png")


@app.route("/admin/stats", methods=["GET"])
@require_admin
def admin_stats():
    """재고 현황 (인증 필요)"""
    stats = db.get_inventory_stats()
    return jsonify({"status": "ok", "inventory": stats})


@app.route("/admin/orders", methods=["GET"])
@require_admin
def admin_orders():
    """주문 목록 조회 API"""
    limit = int(request.args.get("limit", 20))
    offset = int(request.args.get("offset", 0))
    status = request.args.get("status") or None
    search = request.args.get("search") or None
    orders, total = db.get_all_orders(limit=limit, offset=offset, status=status, search=search)
    return jsonify({"orders": orders, "total": total, "limit": limit, "offset": offset})


@app.route("/admin/esims", methods=["GET"])
@require_admin
def admin_esims():
    """eSIM 재고 목록 조회 API"""
    limit = int(request.args.get("limit", 20))
    offset = int(request.args.get("offset", 0))
    status = request.args.get("status") or None
    search = request.args.get("search") or None
    plan = request.args.get("plan") or None
    esims, total = db.get_all_esims(limit=limit, offset=offset, status=status, search=search, plan=plan)
    return jsonify({"esims": esims, "total": total, "limit": limit, "offset": offset})


# ========== Event Handlers ==========

@handler.add(FollowEvent)
def handle_follow(event):
    """친구 추가 시 환영 메시지"""
    reply(event, msg.WELCOME_MESSAGE)


@handler.add(MessageEvent, message=TextMessageContent)
def handle_text_message(event):
    """텍스트 메시지 처리"""
    user_id = event.source.user_id
    text = event.message.text.strip()

    # 세션 가져오기
    session = db.get_or_create_session(user_id)
    current_step = session["current_step"]

    # ── 리치 메뉴 버튼 액션 처리 ──
    if text == msg.ACTION_QR_ISSUE:
        db.update_session(user_id, "awaiting_order_number_qr")
        reply(event, msg.ASK_ORDER_NUMBER)
        return

    if text == msg.ACTION_HOW_TO_USE:
        db.update_session(user_id, "idle")
        reply(event, msg.HOW_TO_USE_GUIDE)
        return

    if text == msg.ACTION_KOREA_GUIDE:
        db.update_session(user_id, "idle")
        reply(event, msg.KOREA_ESIM_VOICE_GUIDE)
        return

    if text == msg.ACTION_CONTACT:
        db.update_session(user_id, "idle")
        reply(event, msg.CONTACT_MESSAGE)
        return

    # ── QR코드 발행 플로우: 주문번호 입력 대기 중 ──
    if current_step == "awaiting_order_number_qr":
        handle_qr_issue(event, user_id, text)
        return

    # ── 기본 메시지 (그 외) ──
    # "こんにちは", "hi" 등 인사에 대응
    greetings = ["こんにちは", "こんばんは", "おはよう", "hi", "hello", "ハロー"]
    if any(g in text.lower() for g in greetings):
        reply(event, msg.WELCOME_MESSAGE)
        return

    reply(event, msg.UNKNOWN_MESSAGE)


# ========== Business Logic ==========

def handle_qr_issue(event, user_id, text):
    """QR코드 발행 처리"""
    # 주문번호 형식 검증
    if not AMAZON_ORDER_PATTERN.match(text):
        reply(event, msg.ORDER_NOT_FOUND)
        return

    # 주문 조회
    order = db.get_order_by_amazon_id(text)
    if not order:
        order = db.get_order_by_id(text)

    if not order:
        reply(event, msg.ORDER_NOT_FOUND)
        db.update_session(user_id, "idle")
        return

    # 이미 전달 완료된 경우
    if order["status"] == "delivered":
        reply(event, msg.ORDER_ALREADY_DELIVERED)
        db.update_session(user_id, "idle")
        return

    # eSIM이 매칭되어 있는지 확인
    if not order.get("qr_code_data"):
        reply(event, "⏳ eSIMの準備中です。しばらくお待ちください。")
        db.update_session(user_id, "idle")
        return

    # QR코드 전달
    _deliver_esim_qr(event, user_id, order)


def _deliver_esim_qr(event, user_id, order):
    """인증 완료 후 QR코드 + eSIM 정보 전달"""
    delivery_text = msg.ESIM_DELIVERED.format(
        order_id=order.get("amazon_order_id") or order["order_id"],
        plan_name=order.get("plan_name", "—"),
        data_amount=order.get("data_amount", "無制限"),
        phone_number=order.get("phone_number", "—"),
        rental_number=order.get("rental_number", "—"),
        sm_dp_address=order.get("sm_dp_address", "—"),
        activation_code=order.get("activation_code", "—"),
    )

    messages = []

    # QR코드 이미지 전송
    qr_code_data = order.get("qr_code_data")
    if qr_code_data:
        if qr_code_data.startswith("http"):
            messages.append(
                ImageMessage(
                    original_content_url=qr_code_data,
                    preview_image_url=qr_code_data,
                )
            )
        else:
            # LPA 데이터 → 서버에서 QR 이미지 생성
            token = db.generate_qr_token(order["order_id"])
            base_url = request.url_root.rstrip("/")
            if base_url.startswith("http://"):
                base_url = "https://" + base_url[7:]
            qr_url = f"{base_url}/qr/{token}"
            logger.info(f"QR image URL: {qr_url}")
            messages.append(
                ImageMessage(
                    original_content_url=qr_url,
                    preview_image_url=qr_url,
                )
            )

    messages.append(TextMessage(text=delivery_text))

    # reply 먼저 보내고 성공하면 delivered 처리
    try:
        reply_messages(event, messages)
        db.mark_esim_delivered(order["order_id"])
    except Exception as e:
        logger.error(f"Reply failed, sending text only: {e}")
        try:
            reply(event, delivery_text)
            db.mark_esim_delivered(order["order_id"])
        except Exception as e2:
            logger.error(f"Text reply also failed: {e2}")

    db.update_session(user_id, "idle")


# ========== Reply Helpers ==========

def reply(event, text):
    """단일 텍스트 메시지 응답"""
    with ApiClient(line_config) as api_client:
        api = MessagingApi(api_client)
        api.reply_message(
            ReplyMessageRequest(
                reply_token=event.reply_token,
                messages=[TextMessage(text=text)],
            )
        )


def reply_messages(event, messages):
    """복수 메시지 응답"""
    with ApiClient(line_config) as api_client:
        api = MessagingApi(api_client)
        api.reply_message(
            ReplyMessageRequest(
                reply_token=event.reply_token,
                messages=messages,
            )
        )


# ========== Admin: Order Sync ==========

@app.route("/admin/sync-orders", methods=["GET"])
@require_admin
def sync_orders():
    """Amazon 주문 수동 동기화"""
    hours = request.args.get("hours", 24, type=int)
    result = amazon_api.sync_orders(hours_back=hours)
    return jsonify(result)


@app.route("/admin/create-test-order", methods=["POST"])
@require_admin
def create_test_order():
    """テスト注文作成 + eSIM自動マッチング"""
    import random
    # テスト用Amazon注文番号生成 (250-xxxxxxx-xxxxxxx)
    order_num = f"250-{random.randint(1000000,9999999)}-{random.randint(1000000,9999999)}"

    created = db.create_order(
        order_id=order_num,
        amazon_order_id=order_num,
        buyer_email="test@example.com",
    )
    if not created:
        return jsonify({"error": "Order creation failed"}), 400

    # 미할당 eSIM 자동 매칭
    esim = db.get_unassigned_esim()
    if esim:
        db.assign_esim_to_order(esim["id"], order_num)
        return jsonify({
            "success": True,
            "order_id": order_num,
            "matched_iccid": esim["iccid"],
            "message": f"LINEで「{order_num}」を入力してテストしてください"
        })
    else:
        return jsonify({
            "success": True,
            "order_id": order_num,
            "matched_iccid": None,
            "message": "注文作成済み（未マッチ：eSIM在庫なし）"
        })


@app.route("/admin/reset-test", methods=["POST"])
@require_admin
def reset_test_data():
    """テストデータリセット（全注文削除 + eSIM在庫をunassignedに戻す）"""
    conn = db.get_connection()
    conn.execute("DELETE FROM order_mapping")
    conn.execute("UPDATE esim_inventory SET status = 'unassigned', updated_at = CURRENT_TIMESTAMP")
    conn.commit()
    conn.close()
    return jsonify({"success": True, "message": "全注文を削除し、eSIM在庫をリセットしました"})


# ========== Admin: 업로드 페이지 ==========

UPLOAD_PAGE_HTML = """<!DOCTYPE html>
<html lang="ja">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>eSIM Admin</title>
<style>
* { box-sizing: border-box; margin: 0; padding: 0; }
body { font-family: -apple-system, sans-serif; background: #f5f5f5; padding: 20px; }
.card { background: #fff; border-radius: 12px; padding: 24px; margin-bottom: 16px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.1); max-width: 600px; margin: 16px auto; }
h1 { font-size: 20px; margin-bottom: 16px; color: #333; }
h2 { font-size: 16px; margin-bottom: 12px; color: #555; }
.upload-area { border: 2px dashed #ccc; border-radius: 8px; padding: 32px;
               text-align: center; cursor: pointer; transition: border-color 0.2s; }
.upload-area:hover { border-color: #2196F3; }
.upload-area.dragover { border-color: #2196F3; background: #E3F2FD; }
input[type="file"] { display: none; }
.btn { background: #2196F3; color: #fff; border: none; padding: 12px 24px;
       border-radius: 8px; font-size: 15px; cursor: pointer; width: 100%; margin-top: 12px; }
.btn:hover { background: #1976D2; }
.btn:disabled { background: #ccc; cursor: not-allowed; }
.btn-danger { background: #f44336; }
.btn-danger:hover { background: #d32f2f; }
.result { margin-top: 16px; padding: 16px; border-radius: 8px; font-size: 14px; display: none; }
.result.success { background: #E8F5E9; color: #2E7D32; }
.result.error { background: #FFEBEE; color: #C62828; }
.stats { display: grid; grid-template-columns: repeat(2, 1fr); gap: 8px; }
.stat { background: #f9f9f9; padding: 12px; border-radius: 8px; text-align: center; }
.stat-num { font-size: 24px; font-weight: bold; color: #2196F3; }
.stat-label { font-size: 12px; color: #888; margin-top: 4px; }
.loading { display: none; text-align: center; margin-top: 12px; color: #888; }
</style>
</head>
<body>
<div class="card">
  <h1>📱 eSIM 在庫管理</h1>
  <div id="stats"><div class="loading" style="display:block">読み込み中...</div></div>
</div>

<div class="card">
  <h2>📤 在庫アップロード</h2>
  <div class="upload-area" id="dropZone" onclick="document.getElementById('fileInput').click()">
    <p style="font-size:32px;margin-bottom:8px">📁</p>
    <p>Excelファイルをドラッグ＆ドロップ</p>
    <p style="color:#aaa;font-size:13px;margin-top:4px">またはクリックして選択</p>
  </div>
  <input type="file" id="fileInput" accept=".xlsx,.xls">
  <div id="fileName" style="margin-top:8px;color:#555;font-size:14px"></div>
  <button class="btn" id="uploadBtn" onclick="upload()" disabled>アップロード</button>
  <div class="loading" id="uploadLoading">⏳ アップロード中...</div>
  <div class="result" id="uploadResult"></div>
</div>

<div class="card">
  <h2>🧪 テスト注文作成</h2>
  <p style="font-size:13px;color:#888;margin-bottom:12px">テスト用の注文番号を生成し、eSIMを自動マッチングします</p>
  <button class="btn" style="background:#FF9800" onclick="createTestOrder()">テスト注文を作成</button>
  <div class="loading" id="testLoading">⏳ 作成中...</div>
  <div class="result" id="testResult"></div>
</div>

<div class="card">
  <h2>🔄 Amazon注文同期</h2>
  <button class="btn" onclick="syncOrders()">注文を同期する</button>
  <div class="loading" id="syncLoading">⏳ 同期中...</div>
  <div class="result" id="syncResult"></div>
</div>

<div class="card">
  <h2>📦 eSIM在庫一覧</h2>
  <div style="display:flex;gap:8px;margin-bottom:12px">
    <input type="text" id="esimSearch" placeholder="ICCID・電話番号で検索"
      style="flex:1;padding:8px 12px;border:1px solid #ddd;border-radius:6px;font-size:14px">
    <select id="esimStatus" style="padding:8px;border:1px solid #ddd;border-radius:6px;font-size:14px">
      <option value="">全て</option>
      <option value="unassigned">未割当</option>
      <option value="assigned">割当済</option>
      <option value="delivered">配信済</option>
    </select>
    <select id="esimPlan" style="padding:8px;border:1px solid #ddd;border-radius:6px;font-size:14px">
      <option value="">全プラン</option>
      <option value="1DAYS">1日</option><option value="3DAYS">3日</option>
      <option value="5DAYS">5日</option><option value="7DAYS">7日</option>
      <option value="10DAYS">10日</option><option value="15DAYS">15日</option>
      <option value="20DAYS">20日</option><option value="30DAYS">30日</option>
      <option value="60DAYS">60日</option><option value="90DAYS">90日</option>
    </select>
    <button onclick="loadEsims(0)" style="padding:8px 16px;background:#4CAF50;color:#fff;border:none;border-radius:6px;cursor:pointer">検索</button>
  </div>
  <div id="esimList" style="font-size:13px">読み込み中...</div>
  <div id="esimPaging" style="text-align:center;margin-top:12px"></div>
</div>

<div class="card">
  <h2>📋 注文一覧</h2>
  <div style="display:flex;gap:8px;margin-bottom:12px">
    <input type="text" id="searchInput" placeholder="注文番号・ICCIDで検索"
      style="flex:1;padding:8px 12px;border:1px solid #ddd;border-radius:6px;font-size:14px">
    <select id="statusFilter" style="padding:8px;border:1px solid #ddd;border-radius:6px;font-size:14px">
      <option value="">全て</option>
      <option value="pending">⏳ 準備中</option>
      <option value="matched">✅ マッチ済</option>
      <option value="delivered">📦 配信済</option>
    </select>
    <button onclick="loadOrders()" style="padding:8px 16px;background:#2196F3;color:#fff;border:none;border-radius:6px;cursor:pointer">検索</button>
  </div>
  <div id="orderList" style="font-size:13px">読み込み中...</div>
  <div id="orderPaging" style="text-align:center;margin-top:12px"></div>
</div>

<div class="card">
  <h2>🗑️ テストデータリセット</h2>
  <p style="font-size:13px;color:#888;margin-bottom:12px">全注文を削除し、eSIM在庫を未割当に戻します</p>
  <button class="btn btn-danger" onclick="resetTest()">リセット実行</button>
  <div class="result" id="resetResult"></div>
</div>

<script>
const KEY = new URLSearchParams(location.search).get('key') || '';
const BASE = location.origin;

// 統計
async function loadStats() {
  try {
    const r = await fetch(BASE + '/admin/stats?key=' + KEY);
    const d = await r.json();
    if (d.error) { document.getElementById('stats').innerHTML = '<p style="color:red">' + d.error + '</p>'; return; }
    const s = d.inventory;
    document.getElementById('stats').innerHTML =
      '<div class="stats">' +
      '<div class="stat"><div class="stat-num">' + s.total + '</div><div class="stat-label">総数</div></div>' +
      '<div class="stat"><div class="stat-num">' + s.unassigned + '</div><div class="stat-label">未割当</div></div>' +
      '<div class="stat"><div class="stat-num">' + s.assigned + '</div><div class="stat-label">割当済</div></div>' +
      '<div class="stat"><div class="stat-num">' + s.delivered + '</div><div class="stat-label">配信済</div></div>' +
      '</div>';
  } catch(e) { document.getElementById('stats').innerHTML = '<p style="color:red">読み込み失敗</p>'; }
}

// ファイル選択
const fileInput = document.getElementById('fileInput');
const dropZone = document.getElementById('dropZone');
let selectedFile = null;

fileInput.addEventListener('change', (e) => { selectFile(e.target.files[0]); });
dropZone.addEventListener('dragover', (e) => { e.preventDefault(); dropZone.classList.add('dragover'); });
dropZone.addEventListener('dragleave', () => { dropZone.classList.remove('dragover'); });
dropZone.addEventListener('drop', (e) => { e.preventDefault(); dropZone.classList.remove('dragover'); selectFile(e.dataTransfer.files[0]); });

function selectFile(f) {
  if (!f) return;
  selectedFile = f;
  document.getElementById('fileName').textContent = '📎 ' + f.name;
  document.getElementById('uploadBtn').disabled = false;
}

// アップロード
async function upload() {
  if (!selectedFile) return;
  const btn = document.getElementById('uploadBtn');
  const loading = document.getElementById('uploadLoading');
  const result = document.getElementById('uploadResult');
  btn.disabled = true; loading.style.display = 'block'; result.style.display = 'none';

  try {
    const fd = new FormData();
    fd.append('file', selectedFile);
    const r = await fetch(BASE + '/admin/upload-inventory?key=' + KEY, { method: 'POST', body: fd });
    const d = await r.json();
    if (d.success) {
      result.className = 'result success';
      result.innerHTML = '✅ 追加: ' + d.added + '件 / スキップ(重複): ' + d.skipped_duplicates + '件' +
        (d.errors.length ? '<br>⚠️ エラー: ' + d.errors.join(', ') : '');
      loadStats();
    } else {
      result.className = 'result error';
      result.textContent = '❌ ' + (d.error || 'アップロード失敗');
    }
  } catch(e) {
    result.className = 'result error';
    result.textContent = '❌ 通信エラー: ' + e.message;
  }
  result.style.display = 'block'; loading.style.display = 'none'; btn.disabled = false;
}

// 注文同期
async function syncOrders() {
  const loading = document.getElementById('syncLoading');
  const result = document.getElementById('syncResult');
  loading.style.display = 'block'; result.style.display = 'none';
  try {
    const r = await fetch(BASE + '/admin/sync-orders?key=' + KEY);
    const d = await r.json();
    result.className = 'result success';
    result.innerHTML = '新規: ' + d.new + ' / マッチ: ' + d.matched + ' / スキップ: ' + d.skipped + ' / エラー: ' + d.errors;
  } catch(e) {
    result.className = 'result error';
    result.textContent = '❌ ' + e.message;
  }
  result.style.display = 'block'; loading.style.display = 'none';
}

// テスト注文作成
async function createTestOrder() {
  const loading = document.getElementById('testLoading');
  const result = document.getElementById('testResult');
  loading.style.display = 'block'; result.style.display = 'none';
  try {
    const r = await fetch(BASE + '/admin/create-test-order?key=' + KEY, { method: 'POST' });
    const d = await r.json();
    if (d.success) {
      result.className = 'result success';
      result.innerHTML = '✅ 注文番号: <strong>' + d.order_id + '</strong><br>' +
        (d.matched_iccid ? 'マッチ済みICCID: ' + d.matched_iccid : '⚠️ eSIM在庫なし') +
        '<br><br>📱 LINEで上の注文番号を入力してテストしてください';
      loadStats();
    } else {
      result.className = 'result error';
      result.textContent = '❌ ' + (d.error || '作成失敗');
    }
  } catch(e) {
    result.className = 'result error';
    result.textContent = '❌ ' + e.message;
  }
  result.style.display = 'block'; loading.style.display = 'none';
}

// リセット
async function resetTest() {
  if (!confirm('本当にリセットしますか？全注文が削除されます。')) return;
  const result = document.getElementById('resetResult');
  try {
    const r = await fetch(BASE + '/admin/reset-test?key=' + KEY, { method: 'POST' });
    const d = await r.json();
    result.className = 'result success';
    result.textContent = '✅ ' + d.message;
    loadStats();
  } catch(e) {
    result.className = 'result error';
    result.textContent = '❌ ' + e.message;
  }
  result.style.display = 'block';
}

// eSIM在庫一覧
let esimPage = 0;

async function loadEsims(page) {
  if (page !== undefined) esimPage = page;
  const search = document.getElementById('esimSearch').value.trim();
  const status = document.getElementById('esimStatus').value;
  const plan = document.getElementById('esimPlan').value;
  const offset = esimPage * PAGE_SIZE;
  const list = document.getElementById('esimList');
  const paging = document.getElementById('esimPaging');
  list.innerHTML = '読み込み中...';
  try {
    let url = BASE + '/admin/esims?key=' + KEY + '&limit=' + PAGE_SIZE + '&offset=' + offset;
    if (search) url += '&search=' + encodeURIComponent(search);
    if (status) url += '&status=' + status;
    if (plan) url += '&plan=' + plan;
    const r = await fetch(url);
    const d = await r.json();
    if (!d.esims.length) { list.innerHTML = '<p style="color:#888;text-align:center">データがありません</p>'; paging.innerHTML=''; return; }
    const sMap = {unassigned:'🟢未割当', assigned:'🟡割当済', delivered:'🔵配信済'};
    let html = '<table style="width:100%;border-collapse:collapse">' +
      '<tr style="background:#f0f0f0"><th style="padding:6px;text-align:left">ICCID</th><th>電話番号</th><th>プラン</th><th>ステータス</th></tr>';
    d.esims.forEach(e => {
      html += '<tr style="border-bottom:1px solid #eee">' +
        '<td style="padding:8px 4px;font-size:11px">' + (e.iccid||'-') + '</td>' +
        '<td style="padding:8px 4px;font-size:12px">' + (e.phone_number||'-') + '</td>' +
        '<td style="padding:8px 4px;font-size:12px">' + (e.plan_name||'-') + '</td>' +
        '<td style="padding:8px 4px;font-size:12px">' + (sMap[e.status]||e.status) + '</td></tr>';
    });
    html += '</table>';
    list.innerHTML = html;
    const totalPages = Math.ceil(d.total / PAGE_SIZE);
    let ph = '';
    if (esimPage > 0) ph += '<button onclick="loadEsims('+(esimPage-1)+')" style="margin:0 4px;padding:4px 12px;cursor:pointer">◀ 前</button>';
    ph += ' ' + (esimPage+1) + ' / ' + totalPages + ' (' + d.total + '件) ';
    if (esimPage < totalPages - 1) ph += '<button onclick="loadEsims('+(esimPage+1)+')" style="margin:0 4px;padding:4px 12px;cursor:pointer">次 ▶</button>';
    paging.innerHTML = ph;
  } catch(e) { list.innerHTML = '<p style="color:red">読み込み失敗: '+e.message+'</p>'; }
}

document.getElementById('esimSearch').addEventListener('keydown', (e) => { if(e.key==='Enter') loadEsims(0); });

// 注文一覧
let orderPage = 0;
const PAGE_SIZE = 20;

async function loadOrders(page) {
  if (page !== undefined) orderPage = page;
  const search = document.getElementById('searchInput').value.trim();
  const status = document.getElementById('statusFilter').value;
  const offset = orderPage * PAGE_SIZE;
  const list = document.getElementById('orderList');
  const paging = document.getElementById('orderPaging');
  list.innerHTML = '読み込み中...';
  try {
    let url = BASE + '/admin/orders?key=' + KEY + '&limit=' + PAGE_SIZE + '&offset=' + offset;
    if (search) url += '&search=' + encodeURIComponent(search);
    if (status) url += '&status=' + status;
    const r = await fetch(url);
    const d = await r.json();
    if (!d.orders.length) { list.innerHTML = '<p style="color:#888;text-align:center">注文がありません</p>'; paging.innerHTML=''; return; }
    const statusMap = {pending:'⏳準備中', matched:'✅マッチ済', delivered:'📦配信済'};
    let html = '<table style="width:100%;border-collapse:collapse">' +
      '<tr style="background:#f0f0f0"><th style="p:6px;text-align:left">注文番号</th><th>ICCID</th><th>プラン</th><th>ステータス</th><th>日時</th></tr>';
    d.orders.forEach(o => {
      html += '<tr style="border-bottom:1px solid #eee">' +
        '<td style="padding:8px 4px;font-size:12px">' + (o.order_id||'-') + '</td>' +
        '<td style="padding:8px 4px;font-size:12px">' + (o.iccid||'-') + '</td>' +
        '<td style="padding:8px 4px;font-size:12px">' + (o.plan_name||'-') + '</td>' +
        '<td style="padding:8px 4px;font-size:12px">' + (statusMap[o.status]||o.status) + '</td>' +
        '<td style="padding:8px 4px;font-size:12px">' + (o.created_at||'').slice(0,16) + '</td></tr>';
    });
    html += '</table>';
    list.innerHTML = html;
    // paging
    const totalPages = Math.ceil(d.total / PAGE_SIZE);
    let ph = '';
    if (orderPage > 0) ph += '<button onclick="loadOrders('+(orderPage-1)+')" style="margin:0 4px;padding:4px 12px;cursor:pointer">◀ 前</button>';
    ph += ' ' + (orderPage+1) + ' / ' + totalPages + ' ';
    if (orderPage < totalPages - 1) ph += '<button onclick="loadOrders('+(orderPage+1)+')" style="margin:0 4px;padding:4px 12px;cursor:pointer">次 ▶</button>';
    paging.innerHTML = ph;
  } catch(e) { list.innerHTML = '<p style="color:red">読み込み失敗: '+e.message+'</p>'; }
}

document.getElementById('searchInput').addEventListener('keydown', (e) => { if(e.key==='Enter') loadOrders(0); });

loadStats();
loadEsims(0);
loadOrders(0);
</script>
</body>
</html>"""


@app.route("/admin/dashboard", methods=["GET"])
@require_admin
def admin_dashboard():
    """Admin 管理画面"""
    return UPLOAD_PAGE_HTML


# ========== Admin: eSIM 재고 업로드 ==========

def _parse_plan_days(plan_str):
    """Plan 문자열에서 유효기간(일) 파싱. 예: '1DAYS' → 1, '7DAYS' → 7, '30days' → 30"""
    if not plan_str:
        return 0
    plan_str = str(plan_str).strip().upper()
    m = re.search(r'(\d+)\s*DAY', plan_str)
    return int(m.group(1)) if m else 0


# 원본 Excel 컬럼명 → DB 컬럼명 매핑 (대소문자 무시)
COLUMN_MAP = {
    "iccid": "iccid",
    "sm-dp+address": "sm_dp_address",
    "sm_dp_address": "sm_dp_address",
    "smdp+address": "sm_dp_address",
    "activation code": "activation_code",
    "activation_code": "activation_code",
    "qr_code": "qr_code_data",
    "qr_code_data": "qr_code_data",
    "plan": "plan_name",
    "plan_name": "plan_name",
    "data_amount": "data_amount",
    "validity_days": "validity_days",
    "country": "country",
    "렌탈관리번호": "rental_number",
    "rental_number": "rental_number",
    "phone no.": "phone_number",
    "phone_number": "phone_number",
    "phone no": "phone_number",
}


@app.route("/admin/upload-inventory", methods=["POST"])
@require_admin
def upload_inventory():
    """Excel 파일로 eSIM 재고 업로드

    원본 형식 (SM-DP+Address, Activation Code, QR_CODE 등) 자동 인식
    """
    import openpyxl

    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file uploaded"}), 400

    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)

        added = 0
        skipped = 0
        errors = []
        sheets_processed = []

        # 모든 시트 순회
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # 헤더 읽기 + 컬럼 매핑
            first_row = next(ws.iter_rows(min_row=1, max_row=1), None)
            if not first_row:
                continue

            raw_headers = [str(cell.value).strip() if cell.value else "" for cell in first_row]
            mapped_headers = []
            for h in raw_headers:
                key = h.lower().replace("+", "+")
                mapped = COLUMN_MAP.get(key, key)
                mapped_headers.append(mapped)

            # 필수 컬럼 확인
            required = {"iccid", "activation_code", "qr_code_data"}
            found = set(mapped_headers)
            if not required.issubset(found):
                errors.append(f"Sheet '{sheet_name}': missing columns {required - found}")
                continue

            sheet_added = 0
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                data = dict(zip(mapped_headers, row))
                iccid = str(data.get("iccid", "")).strip()
                if not iccid:
                    continue

                # Plan에서 validity_days 자동 파싱 (기본값: 30일)
                plan_name = str(data.get("plan_name", "") or "").strip()
                validity_days = data.get("validity_days")
                if not validity_days:
                    validity_days = _parse_plan_days(plan_name) or 30
                else:
                    validity_days = int(validity_days)

                # データ量 기본값: 無制限
                data_amount = str(data.get("data_amount", "") or "").strip() or "無制限"

                try:
                    ok = db.add_esim(
                        iccid=iccid,
                        rental_number=str(data.get("rental_number", "") or "").strip() or None,
                        phone_number=str(data.get("phone_number", "") or "").strip() or None,
                        sm_dp_address=str(data.get("sm_dp_address", "") or "").strip(),
                        activation_code=str(data.get("activation_code", "") or "").strip(),
                        qr_code_data=str(data.get("qr_code_data", "") or "").strip(),
                        plan_name=plan_name,
                        data_amount=data_amount,
                        validity_days=validity_days,
                        country=str(data.get("country", "JP") or "JP").strip(),
                    )
                    if ok:
                        added += 1
                        sheet_added += 1
                    else:
                        skipped += 1  # 중복 ICCID
                except Exception as e:
                    errors.append(f"Sheet '{sheet_name}' Row {row_idx}: {str(e)}")

            sheets_processed.append({"sheet": sheet_name, "added": sheet_added})

        wb.close()
        return jsonify({
            "success": True,
            "added": added,
            "skipped_duplicates": skipped,
            "sheets_processed": sheets_processed,
            "errors": errors,
        })

    except Exception as e:
        return jsonify({"error": f"File processing failed: {str(e)}"}), 400


# ========== Admin: Rich Menu Setup ==========

def _load_fonts():
    """CJK 폰트 로드"""
    font_paths = [
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Bold.ttc",
    ]
    for fp in font_paths:
        try:
            return {
                "large": ImageFont.truetype(fp, 60),
                "small": ImageFont.truetype(fp, 36),
                "icon": ImageFont.truetype(fp, 40),
                "bar": ImageFont.truetype(fp, 48),
            }
        except Exception:
            continue
    df = ImageFont.load_default()
    return {"large": df, "small": df, "icon": df, "bar": df}


def _hex_rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


def _build_main_image(fonts):
    """메인 리치 메뉴 이미지 (2x2: 3기능 + チャット入力)"""
    W, H = 2500, 843
    CW, CH = W // 2, H // 2
    buttons = [
        {"label": "QRコード発行", "bg": "#2196F3", "desc": "eSIMを受け取る", "icon": "QR"},
        {"label": "使い方ガイド", "bg": "#4CAF50", "desc": "設定方法を確認", "icon": "?"},
        {"label": "音声・SMS案内", "bg": "#FF9800", "desc": "音声通話の利用方法", "icon": "📞"},
        {"label": "チャット入力", "bg": "#607D8B", "desc": "キーボードで入力", "icon": "Aa"},
    ]
    img = Image.new("RGB", (W, H), "#FFFFFF")
    draw = ImageDraw.Draw(img)

    for i, btn in enumerate(buttons):
        col, row = i % 2, i // 2
        x0, y0 = col * CW, row * CH
        x1, y1 = x0 + CW, y0 + CH
        cx = x0 + CW // 2

        draw.rectangle([x0, y0, x1, y1], fill=_hex_rgb(btn["bg"]))
        draw.rectangle([x0, y0, x1, y1], outline="#FFFFFF", width=3)

        icon_cy = y0 + CH // 2 - 60
        r = 45
        draw.ellipse([cx-r, icon_cy-r, cx+r, icon_cy+r],
                     fill=None, outline="#FFFFFF", width=3)
        draw.text((cx, icon_cy), btn["icon"], fill="#FFFFFF",
                  font=fonts["icon"], anchor="mm")

        label_cy = y0 + CH // 2 + 25
        draw.text((cx, label_cy), btn["label"], fill="#FFFFFF",
                  font=fonts["large"], anchor="mm")
        draw.text((cx, label_cy + 55), btn["desc"], fill="#FFFFFFCC",
                  font=fonts["small"], anchor="mm")

    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _build_keyboard_image(fonts):
    """키보드 모드 리치 메뉴 이미지 (메뉴로 돌아가기 바)"""
    W, H = 2500, 843
    img = Image.new("RGB", (W, H), "#37474F")
    draw = ImageDraw.Draw(img)
    draw.text((W // 2, H // 2), "▼ メニューを表示", fill="#FFFFFF",
              font=fonts["bar"], anchor="mm")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _line_api(method, url, token, **kwargs):
    """LINE API 호출 헬퍼"""
    headers = kwargs.pop("headers", {})
    headers["Authorization"] = f"Bearer {token}"
    fn = getattr(http_requests, method)
    return fn(url, headers=headers, **kwargs)


@app.route("/admin/setup-richmenu", methods=["GET"])
@require_admin
def setup_rich_menu():
    """リッチメニュー (メイン ↔ キーボード切替) をワンクリック登録"""
    token = config.LINE_CHANNEL_ACCESS_TOKEN
    if not token:
        return jsonify({"error": "LINE_CHANNEL_ACCESS_TOKEN not set"}), 500

    results = {}

    # ── Step 0: 기존 리치 메뉴 & 에일리어스 전부 삭제 ──
    try:
        aliases = _line_api("get", "https://api.line.me/v2/bot/richmenu/alias/list",
                            token).json().get("aliases", [])
        for a in aliases:
            _line_api("delete",
                      f"https://api.line.me/v2/bot/richmenu/alias/{a['richMenuAliasId']}",
                      token)
    except Exception:
        pass

    try:
        old = _line_api("get", "https://api.line.me/v2/bot/richmenu/list",
                        token).json().get("richmenus", [])
        for m in old:
            _line_api("delete",
                      f"https://api.line.me/v2/bot/richmenu/{m['richMenuId']}", token)
    except Exception:
        pass

    fonts = _load_fonts()

    # ── Step 1: 메인 메뉴 (Menu A) 생성 ──
    W, H = 2500, 843
    CW, CH = W // 2, H // 2

    menu_a_data = {
        "size": {"width": W, "height": H},
        "selected": True,
        "name": "メインメニュー",
        "chatBarText": "メニューを開く",
        "areas": [
            {"bounds": {"x": 0,  "y": 0,  "width": CW, "height": CH},
             "action": {"type": "message", "label": "QRコード発行",
                        "text": "QRコード発行"}},
            {"bounds": {"x": CW, "y": 0,  "width": CW, "height": CH},
             "action": {"type": "message", "label": "使い方ガイド",
                        "text": "使い方ガイド"}},
            {"bounds": {"x": 0,  "y": CH, "width": CW, "height": CH},
             "action": {"type": "message", "label": "音声・SMS案内",
                        "text": "音声・SMS案内"}},
            {"bounds": {"x": CW, "y": CH, "width": CW, "height": CH},
             "action": {"type": "richmenuswitch",
                        "richMenuAliasId": "richmenu-keyboard",
                        "data": "switch-to-keyboard"}},
        ],
    }

    r = _line_api("post", "https://api.line.me/v2/bot/richmenu", token,
                  headers={"Content-Type": "application/json"}, json=menu_a_data)
    if r.status_code != 200:
        return jsonify({"error": "Menu A create failed", "detail": r.text}), 500
    menu_a_id = r.json()["richMenuId"]
    results["menuA"] = menu_a_id

    # Menu A 이미지 업로드
    r = _line_api("post",
                  f"https://api-data.line.me/v2/bot/richmenu/{menu_a_id}/content",
                  token, headers={"Content-Type": "image/png"},
                  data=_build_main_image(fonts))
    if r.status_code != 200:
        return jsonify({"error": "Menu A image failed", "detail": r.text}), 500

    # ── Step 2: 키보드 메뉴 (Menu B) 생성 ──
    menu_b_data = {
        "size": {"width": W, "height": H},
        "selected": False,
        "name": "キーボードモード",
        "chatBarText": "メニューを表示",
        "areas": [
            {"bounds": {"x": 0, "y": 0, "width": W, "height": H},
             "action": {"type": "richmenuswitch",
                        "richMenuAliasId": "richmenu-main",
                        "data": "switch-to-main"}},
        ],
    }

    r = _line_api("post", "https://api.line.me/v2/bot/richmenu", token,
                  headers={"Content-Type": "application/json"}, json=menu_b_data)
    if r.status_code != 200:
        return jsonify({"error": "Menu B create failed", "detail": r.text}), 500
    menu_b_id = r.json()["richMenuId"]
    results["menuB"] = menu_b_id

    # Menu B 이미지 업로드
    r = _line_api("post",
                  f"https://api-data.line.me/v2/bot/richmenu/{menu_b_id}/content",
                  token, headers={"Content-Type": "image/png"},
                  data=_build_keyboard_image(fonts))
    if r.status_code != 200:
        return jsonify({"error": "Menu B image failed", "detail": r.text}), 500

    # ── Step 3: 에일리어스 등록 (richmenuswitch 연결) ──
    for alias_id, mid in [("richmenu-main", menu_a_id),
                          ("richmenu-keyboard", menu_b_id)]:
        r = _line_api("post", "https://api.line.me/v2/bot/richmenu/alias", token,
                      headers={"Content-Type": "application/json"},
                      json={"richMenuAliasId": alias_id, "richMenuId": mid})
        if r.status_code != 200:
            return jsonify({"error": f"Alias {alias_id} failed",
                            "detail": r.text}), 500

    # ── Step 4: Menu A를 기본 리치 메뉴로 설정 ──
    r = _line_api("post",
                  f"https://api.line.me/v2/bot/user/all/richmenu/{menu_a_id}", token)
    if r.status_code != 200:
        return jsonify({"error": "set default failed", "detail": r.text}), 500

    results["success"] = True
    results["message"] = "リッチメニュー(メイン↔キーボード切替)が正常に登録されました！"
    return jsonify(results)


# ========== 주문 자동 동기화 (백그라운드) ==========

import threading

_sync_running = False


def _auto_sync_loop(interval=300):
    """5분(300초)마다 Amazon 주문 자동 동기화"""
    global _sync_running
    if _sync_running:
        return
    _sync_running = True
    logger.info(f"Auto-sync started (interval: {interval}s)")

    def run():
        while True:
            try:
                result = amazon_api.sync_orders(hours_back=1)
                if result["new"] > 0 or result["errors"] > 0:
                    logger.info(f"Auto-sync result: {result}")
            except Exception as e:
                logger.error(f"Auto-sync error: {e}")
            import time as _time
            _time.sleep(interval)

    t = threading.Thread(target=run, daemon=True)
    t.start()


# gunicorn 환경에서도 자동 시작
_auto_sync_loop()


# ========== Startup ==========

if __name__ == "__main__":
    db.init_db()
    print("=" * 50)
    print("  LINE eSIM Bot Started")
    print(f"  Webhook URL: http://{config.HOST}:{config.PORT}/callback")
    print("=" * 50)
    app.run(host=config.HOST, port=config.PORT, debug=True)
