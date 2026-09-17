"""
eSIM在庫エクセルファイルをDBにアップロードするスクリプト

使い方:
  python scripts/upload_inventory.py inventory.xlsx

エクセルファイルの必須列:
  - ICCID          : eSIMのICCID番号
  - SM-DP+ Address : SM-DP+サーバーアドレス
  - Activation Code : アクティベーションコード
  - QR Code Data   : QRコードデータ（LPA:1$...形式）またはQRコード画像URL
  - Plan Name      : プラン名（例：Japan 7日間 3GB）
  - Data Amount    : データ量（例：3GB）
  - Validity Days  : 有効日数（例：7）
  - Country        : 国コード（デフォルト：JP）※省略可
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import database as db


# 엑셀 컬럼명 맵핑 (유연하게 대응)
COLUMN_MAP = {
    "iccid":           ["iccid", "ICCID", "iccid番号"],
    "sm_dp_address":   ["sm-dp+ address", "SM-DP+ Address", "smdp", "sm_dp_address"],
    "activation_code": ["activation code", "Activation Code", "activation_code", "ac"],
    "qr_code_data":    ["qr code data", "QR Code Data", "qr_code_data", "qr", "qrcode"],
    "plan_name":       ["plan name", "Plan Name", "plan_name", "プラン名", "plan"],
    "data_amount":     ["data amount", "Data Amount", "data_amount", "データ量", "data"],
    "validity_days":   ["validity days", "Validity Days", "validity_days", "有効日数", "days"],
    "country":         ["country", "Country", "国", "国コード"],
}


def find_column_index(headers, field_name):
    """헤더 행에서 컬럼 인덱스를 찾기"""
    possible_names = COLUMN_MAP.get(field_name, [field_name])
    for i, h in enumerate(headers):
        if h and str(h).strip() in possible_names:
            return i
    return None


def upload_from_excel(filepath):
    """엑셀 파일에서 eSIM 재고를 DB에 업로드"""
    if not os.path.exists(filepath):
        print(f"❌ ファイルが見つかりません: {filepath}")
        return

    # DB 초기화
    db.init_db()

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("❌ エクセルファイルが空です。")
        return

    # 첫 행 = 헤더
    headers = [str(h).strip() if h else "" for h in rows[0]]

    # 컬럼 인덱스 찾기
    col_indices = {}
    required_fields = ["iccid", "sm_dp_address", "activation_code",
                       "qr_code_data", "plan_name", "data_amount", "validity_days"]

    for field in required_fields + ["country"]:
        idx = find_column_index(headers, field)
        col_indices[field] = idx

    # 필수 컬럼 체크
    missing = [f for f in required_fields if col_indices[f] is None]
    if missing:
        print("❌ 必須列が見つかりません:")
        for m in missing:
            expected = COLUMN_MAP.get(m, [m])
            print(f"   - {m} (期待する列名: {', '.join(expected)})")
        print()
        print("📋 見つかった列:")
        for i, h in enumerate(headers):
            print(f"   [{i}] {h}")
        return

    # 데이터 행 처리
    success = 0
    skipped = 0
    errors = 0

    for row_num, row in enumerate(rows[1:], start=2):
        try:
            iccid = str(row[col_indices["iccid"]]).strip()
            if not iccid or iccid == "None":
                continue

            sm_dp = str(row[col_indices["sm_dp_address"]] or "").strip()
            ac = str(row[col_indices["activation_code"]] or "").strip()
            qr = str(row[col_indices["qr_code_data"]] or "").strip()
            plan = str(row[col_indices["plan_name"]] or "").strip()
            data_amt = str(row[col_indices["data_amount"]] or "").strip()

            validity = row[col_indices["validity_days"]]
            try:
                validity = int(validity)
            except (TypeError, ValueError):
                validity = 0

            country = "JP"
            if col_indices["country"] is not None and row[col_indices["country"]]:
                country = str(row[col_indices["country"]]).strip()

            result = db.add_esim(
                iccid=iccid,
                sm_dp_address=sm_dp,
                activation_code=ac,
                qr_code_data=qr,
                plan_name=plan,
                data_amount=data_amt,
                validity_days=validity,
                country=country,
            )

            if result:
                success += 1
            else:
                skipped += 1  # 중복 ICCID

        except Exception as e:
            errors += 1
            print(f"⚠️  行 {row_num} エラー: {e}")

    wb.close()

    print()
    print("=" * 50)
    print("  アップロード結果")
    print("=" * 50)
    print(f"  ✓ 登録成功: {success} 件")
    print(f"  ⏭️  スキップ（重複）: {skipped} 件")
    print(f"  ❌ エラー: {errors} 件")
    print()

    stats = db.get_inventory_stats()
    print("📊 在庫状況:")
    print(f"  未割当: {stats['unassigned']} 件")
    print(f"  割当済: {stats['assigned']} 件")
    print(f"  配信済: {stats['delivered']} 件")
    print(f"  合計:   {stats['total']} 件")


def create_sample_excel(filepath="sample_inventory.xlsx"):
    """サンプルエクセルファイルを作成"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "eSIM Inventory"

    # 헤더
    headers = ["ICCID", "SM-DP+ Address", "Activation Code",
               "QR Code Data", "Plan Name", "Data Amount",
               "Validity Days", "Country"]
    ws.append(headers)

    # 샘플 데이터 (3행)
    samples = [
        ["8981100000000000001", "smdp.example.com",
         "K2-XXXX-YYYY-ZZZZ", "LPA:1$smdp.example.com$K2-XXXX-YYYY-ZZZZ",
         "Japan 7日間 3GB", "3GB", 7, "JP"],
        ["8981100000000000002", "smdp.example.com",
         "K2-AAAA-BBBB-CCCC", "LPA:1$smdp.example.com$K2-AAAA-BBBB-CCCC",
         "Japan 15日間 5GB", "5GB", 15, "JP"],
        ["8981100000000000003", "smdp.example.com",
         "K2-DDDD-EEEE-FFFF", "LPA:1$smdp.example.com$K2-DDDD-EEEE-FFFF",
         "Japan 30日間 10GB", "10GB", 30, "JP"],
    ]

    for sample in samples:
        ws.append(sample)

    # 열 너비 조정
    for col_idx, header in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + col_idx)].width = max(len(header) + 5, 20)

    wb.save(filepath)
    print(f"✓ サンプルファイルを作成しました: {filepath}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("使い方:")
        print("  python scripts/upload_inventory.py <エクセルファイル>")
        print("  python scripts/upload_inventory.py --sample")
        print()
        print("--sample: サンプルエクセルファイルを生成します")
        sys.exit(1)

    if sys.argv[1] == "--sample":
        output = sys.argv[2] if len(sys.argv) > 2 else "sample_inventory.xlsx"
        create_sample_excel(output)
    else:
        upload_from_excel(sys.argv[1])
